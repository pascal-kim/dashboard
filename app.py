import streamlit as st
import pandas as pd
import plotly.express as px
import re
from datetime import datetime

st.set_page_config(
    page_title="Innovation_Pascal@",
    page_icon="📊",
    layout="wide"
)

# ─────────────────────────────────────────────
# 색상 테마
# ─────────────────────────────────────────────

COLOR_TITLE = "#111D37"
COLOR_BUDGET = "#223D72"
COLOR_BUDGET_SUB = "#345CAD"
COLOR_EXECUTION = "#5F84CF"
COLOR_EXECUTION_SUB = "#9AB1E0"
COLOR_COMPARE = "#7186A8"
COLOR_EXTRA = "#6FA6B8"
COLOR_BACKGROUND = "#D0D9E8"

BLUE_PALETTE = [
    COLOR_BUDGET,
    COLOR_BUDGET_SUB,
    COLOR_EXECUTION,
    COLOR_EXECUTION_SUB,
    COLOR_COMPARE,
    COLOR_EXTRA,
    COLOR_BACKGROUND,
]

st.markdown(
    f"""
    <style>
    h1, h2, h3 {{
        color: {COLOR_TITLE};
    }}

    [data-testid="stMetric"] {{
        background-color: {COLOR_BACKGROUND};
        padding: 14px;
        border-radius: 8px;
    }}

    [data-testid="stMetricLabel"] {{
        color: {COLOR_TITLE};
    }}

    [data-testid="stMetricValue"] {{
        color: {COLOR_TITLE};
        font-weight: 700;
    }}

    [data-testid="stSidebar"] {{
        border-right: 1px solid {COLOR_BACKGROUND};
    }}

    div[data-testid="stDataFrame"] {{
        border: 1px solid {COLOR_BACKGROUND};
        border-radius: 6px;
    }}
    </style>
    """,
    unsafe_allow_html=True
)


# ─────────────────────────────────────────────
# 프린터 최적화 CSS (A4 가로 / 컬러 유지)
# ─────────────────────────────────────────────

st.markdown(
    """
    <style>
    @media print {
        @page {
            size: A4 landscape;
            margin: 10mm;
        }
        html, body, .stApp, [data-testid="stAppViewContainer"] {
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }
        [data-testid="stSidebar"],
        [data-testid="stHeader"],
        [data-testid="stToolbar"],
        [data-testid="stDecoration"],
        [data-testid="stStatusWidget"],
        [data-testid="stFileUploader"],
        [data-testid="stDownloadButton"],
        button {
            display: none !important;
        }
        .main,
        [data-testid="stMain"],
        [data-testid="stMainBlockContainer"] {
            width: 100% !important;
            max-width: 100% !important;
            padding: 0 !important;
            margin: 0 !important;
        }
        h1, h2, h3, h4 {
            break-after: avoid !important;
            page-break-after: avoid !important;
        }
        [data-testid="stPlotlyChart"],
        [data-testid="stDataFrame"],
        [data-testid="stTable"],
        [data-testid="stMetric"],
        .element-container {
            break-inside: avoid !important;
            page-break-inside: avoid !important;
        }
        [data-testid="stMetric"],
        [data-testid="stDataFrame"] {
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }
        table {
            border-collapse: collapse !important;
            width: 100% !important;
        }
        th, td {
            border: 1px solid #D0D9E8 !important;
        }
        a { text-decoration: none !important; }
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.title("혁신지원사업 예산 대시보드")
st.caption(
    "예산 파일과 집행 파일을 함께 올리면 "
    "월별 집행액·예산 대비 집행률을 확인할 수 있습니다."
)


# ─────────────────────────────────────────────
# 공통 함수
# ─────────────────────────────────────────────

def norm(x):
    return re.sub(
        r"\s+",
        "",
        str(x).replace("\n", "")
    ).strip().lower()


def find_col(columns, candidates):
    cmap = {norm(c): c for c in columns}

    for c in candidates:
        if norm(c) in cmap:
            return cmap[norm(c)]

    return None


def read_budget(upload):
    xls = pd.ExcelFile(upload)

    preferred = [
        s for s in xls.sheet_names
        if "예산내역-whatif" in s
    ]

    preferred += [
        s for s in xls.sheet_names
        if "예산내역" in s and s not in preferred
    ]

    preferred += [
        s for s in xls.sheet_names
        if s not in preferred
    ]

    selected = None

    for sh in preferred:
        try:
            test = pd.read_excel(
                xls,
                sheet_name=sh,
                header=1,
                nrows=3
            )

            cn = [norm(c) for c in test.columns]

            if (
                any("부서코드" in c for c in cn)
                and any("자금코드" in c for c in cn)
                and any("최초예산" in c for c in cn)
            ):
                selected = sh
                break

        except Exception:
            pass

    if selected is None:
        selected = xls.sheet_names[0]

    return xls, selected


def area_class(task, area=""):
    t = (
        str(task)
        .strip()
        .upper()
        .replace("[", "")
        .replace("]", "")
    )

    if "사업관리" in t:
        return "사업관리"

    m = re.match(r"^([A-F])[-－]", t)

    if m:
        return "교육영역" if m.group(1) in "ABCDE" else "기타영역"

    a = str(area).strip()

    if a in ("교육", "교육영역"):
        return "교육영역"

    if a in ("기타", "기타영역"):
        return "기타영역"

    if a == "사업관리":
        return "사업관리"

    return "기타/미분류"


def clean_task(x):
    s = str(x).strip().upper()
    s = s.replace("[", "").replace("]", "").replace(" ", "")

    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]

    # SAP 집행 자료의 사업운영은 예산의 사업관리와 연결
    if s == "사업운영":
        s = "사업관리"

    return s


def extract_update_date(filename):
    """파일명에서 YYYYMMDD 형식의 날짜를 찾아 YYYY년 MM월 DD일로 반환."""
    if not filename:
        return None

    # 파일명 중 독립된 8자리 숫자 중 실제 날짜로 유효한 값을 사용
    for match in re.finditer(r"(?<!\d)(\d{8})(?!\d)", str(filename)):
        date_text = match.group(1)
        try:
            parsed = datetime.strptime(date_text, "%Y%m%d")
            return parsed.strftime("%Y년 %m월 %d일")
        except ValueError:
            continue

    return None


def money(x):
    try:
        return f"{float(x):,.0f}원"
    except (ValueError, TypeError):
        return "—"


def format_money_columns(df, columns):
    """표에 표시할 금액을 천 단위 콤마 문자열로 변환."""
    out = df.copy()

    for col in columns:
        if col in out.columns:
            out[col] = out[col].apply(
                lambda x: f"{x:,.0f}"
                if pd.notna(x) and isinstance(x, (int, float))
                else (
                    f"{float(x):,.0f}"
                    if pd.notna(x)
                    and str(x).strip() != ""
                    and is_numeric(x)
                    else x
                )
            )

    return out


def is_numeric(x):
    try:
        float(x)
        return True
    except (ValueError, TypeError):
        return False


def right_align_columns(df, columns):
    """지정한 금액 열을 오른쪽 정렬."""
    cols = [col for col in columns if col in df.columns]

    if not cols:
        return df

    return df.style.set_properties(
        subset=cols,
        **{"text-align": "right"}
    )

def render_nrf_monthly_table(uploaded_file):
    """한국연구재단 월별제출자료를 기존 예산·집행 집계와 분리해 표시."""
    st.subheader("한국연구재단 월별제출자료")

    if uploaded_file is None:
        st.info(
            "한국연구재단 월별제출자료 엑셀 파일을 "
            "업로드하면 이곳에 별도 표로 표시됩니다."
        )
        return

    try:
        import openpyxl

        uploaded_file.seek(0)
        workbook = openpyxl.load_workbook(
            uploaded_file,
            data_only=True
        )

        # 첫 번째 시트 사용
        sheet = workbook[workbook.sheetnames[0]]

        # 병합 셀의 값을 읽을 수 있도록 처리
        merged_values = {}

        for merged_range in sheet.merged_cells.ranges:
            top_left = sheet.cell(
                merged_range.min_row,
                merged_range.min_col
            ).value

            for row_idx in range(
                merged_range.min_row,
                merged_range.max_row + 1
            ):
                for col_idx in range(
                    merged_range.min_col,
                    merged_range.max_col + 1
                ):
                    merged_values[(row_idx, col_idx)] = top_left

        def cell_value(row_idx, col_idx):
            value = sheet.cell(row_idx, col_idx).value

            if value is None:
                return merged_values.get(
                    (row_idx, col_idx),
                    value
                )

            return value

        # 다단계 머리글을 하나의 열 이름으로 결합
        headers = []

        for col_idx in range(1, sheet.max_column + 1):
            if col_idx == 1:
                headers.append("영역")
                continue

            parts = []

            for row_idx in (1, 2, 4, 5):
                value = cell_value(row_idx, col_idx)

                if value is not None:
                    value = str(value).strip()

                    if value and (
                        not parts or parts[-1] != value
                    ):
                        parts.append(value)

            headers.append(
                " / ".join(parts)
                if parts
                else f"열{col_idx}"
            )

        # 데이터 행 읽기
        rows = []

        for row_idx in range(6, sheet.max_row + 1):
            values = [
                sheet.cell(row_idx, col_idx).value
                for col_idx in range(
                    1,
                    sheet.max_column + 1
                )
            ]

            if any(value is not None for value in values):
                rows.append(values)

        table = pd.DataFrame(rows, columns=headers)

        # 빈 열 제거
        table = table.dropna(axis=1, how="all")

        table = table.loc[
            :,
            [
                col for col in table.columns
                if not str(col).startswith("열")
            ]
        ]

        if table.empty:
            st.warning(
                "엑셀 파일에서 표시할 데이터 행을 찾지 못했습니다."
            )
            return

        # 숫자형 열은 천 단위 쉼표 표시
        money_cols = [
            col for col in table.columns
            if col != "영역"
            and pd.api.types.is_numeric_dtype(table[col])
        ]

        st.dataframe(
            table.style.format(
                {col: "{:,.0f}" for col in money_cols},
                na_rep="—"
            ),
            use_container_width=True,
            hide_index=True
        )

        st.caption(
            f"원본 파일: {uploaded_file.name} · "
            "기존 예산·집행 집계와 별도로 표시됩니다."
        )

    except Exception as e:
        st.error(
            f"한국연구재단 월별제출자료를 읽는 중 오류가 발생했습니다: {e}"
        )



# ─────────────────────────────────────────────
# 파일 업로드
# ─────────────────────────────────────────────

st.sidebar.header("파일 업로드")

# ① 예산내역 파일
budget_file = st.sidebar.file_uploader(
    "① 예산내역 엑셀 파일",
    type=["xlsx", "xls"],
    key="budget"
)

# ② 집행내역 파일
execution_file = st.sidebar.file_uploader(
    "② 집행내역 엑셀 파일",
    type=["xlsx", "xls"],
    key="execution"
)

# ③ 한국연구재단 월별제출자료 파일
# 기존 예산 및 집행 집계와 별도로 사용
nrf_monthly_file = st.sidebar.file_uploader(
    "③ 한국연구재단 월별제출자료 엑셀 파일",
    type=["xlsx", "xls"],
    key="nrf_monthly"
)

# ─────────────────────────────────────────────
# 파일명에 포함된 날짜를 기준으로 업데이트일 표시
# 예: 20261029
# ─────────────────────────────────────────────

budget_update_date = (
    extract_update_date(budget_file.name)
    if budget_file is not None
    else None
)

execution_update_date = (
    extract_update_date(execution_file.name)
    if execution_file is not None
    else None
)

# ─────────────────────────────────────────────
# 예산 및 집행 파일 업데이트일
# ─────────────────────────────────────────────

update_col1, update_col2 = st.columns(2)

with update_col1:
    st.markdown("**예산내역 업데이트일:**")

    if budget_file is None:
        st.write("파일 미업로드")
    else:
        st.write(budget_update_date or "날짜 확인 불가")

with update_col2:
    st.markdown("**집행내역 업데이트일:**")

    if execution_file is None:
        st.write("파일 미업로드")
    else:
        st.write(execution_update_date or "날짜 확인 불가")

# ─────────────────────────────────────────────
# 한국연구재단 월별제출자료 업로드 상태
# ─────────────────────────────────────────────

st.sidebar.markdown("---")
st.sidebar.markdown("**한국연구재단 월별제출자료**")

if nrf_monthly_file is None:
    st.sidebar.caption("월별제출자료 파일 미업로드")
else:
    st.sidebar.success(
        f"업로드 완료: {nrf_monthly_file.name}"
    )

# ─────────────────────────────────────────────
# 필수 파일 확인
# ─────────────────────────────────────────────

if not budget_file or not execution_file:

    st.info(
        "예산 파일과 집행 파일을 모두 업로드해 주세요. "
        "예산 파일은 '예산내역' 시트를, 집행 파일은 전표의 "
        "'전기일'과 '금액(현지 통화)'를 사용합니다."
    )

    st.markdown("""
    **대시보드 집계 기준**

    - 예산: 부서명·부서코드·과제번호·자금코드·한국연구재단비목·교내계정과목별 예산 합계
    - 집행: 전기일 기준 월별 집행액(원본 금액의 부호 유지)
    - 영역: `[A-1]`~`[E-3]` 교육영역, `[F-1]`~`[F-3]` 기타영역, `사업관리` 사업관리
    """)

    st.stop()
# ─────────────────────────────────────────────
# 예산 파일 읽기
# ─────────────────────────────────────────────

try:
    bx, default_sheet = read_budget(budget_file)

    bsheet = st.sidebar.selectbox(
        "예산 시트",
        bx.sheet_names,
        index=bx.sheet_names.index(default_sheet)
    )

    bdf = pd.read_excel(
        bx,
        sheet_name=bsheet,
        header=1
    )

except Exception as e:
    st.error(f"예산 파일 읽기 오류: {e}")
    st.stop()


# 예산 열 찾기
bc = {
    "dept": find_col(
        bdf.columns,
        ["부서(팀명)", "부서명", "부서"]
    ),
    "dept_code": find_col(
        bdf.columns,
        ["부서코드"]
    ),
    "task": find_col(
        bdf.columns,
        ["과제명", "과제번호"]
    ),
    "nrf": find_col(
        bdf.columns,
        [
            "한국연구재단비목(드롭박스)",
            "한국연구재단비목",
            "연구재단비목"
        ]
    ),
    "account": find_col(
        bdf.columns,
        ["교내계정과목(드롭박스)", "교내계정과목"]
    ),
    "account_code": find_col(
        bdf.columns,
        ["계정코드"]
    ),
    "fund": find_col(
        bdf.columns,
        ["자금코드"]
    ),
    "budget": find_col(
        bdf.columns,
        ["최초예산", "최종예산", "변경예산", "예산액", "예산"]
    ),
    "area": find_col(
        bdf.columns,
        ["영역1", "사업영역", "영역"]
    )
}

missing_b = [
    k for k in [
        "dept", "dept_code", "task", "nrf",
        "account", "account_code", "fund", "budget"
    ]
    if not bc[k]
]

if missing_b:
    st.error(
        f"예산 시트에서 필요한 열을 찾지 못했습니다: {missing_b}. "
        "상세 예산내역 시트와 제목 행을 확인해 주세요."
    )
    st.write(list(bdf.columns))
    st.stop()


# ─────────────────────────────────────────────
# 집행 파일 읽기
# ─────────────────────────────────────────────

try:
    exls = pd.ExcelFile(execution_file)

    exsheet = st.sidebar.selectbox(
        "집행 시트",
        exls.sheet_names
    )

    raw = pd.read_excel(
        exls,
        sheet_name=exsheet,
        header=None
    )

    header_idx = 0

    for i in range(min(8, len(raw))):
        vals = [norm(v) for v in raw.iloc[i].tolist()]

        if (
            any(v == "전기일" for v in vals)
            and any(
                v in ("금액(현지통화)", "금액(현지 통화)")
                for v in vals
            )
        ):
            header_idx = i
            break

    edf = pd.read_excel(
        exls,
        sheet_name=exsheet,
        header=header_idx
    )

except Exception as e:
    st.error(f"집행 파일 읽기 오류: {e}")
    st.stop()


# 집행 열 찾기
ec = {
    "dept": find_col(
        edf.columns,
        ["지출 부서", "지출부서", "부서명"]
    ),
    "nrf": find_col(
        edf.columns,
        ["연구재단 비목", "연구재단비목", "한국연구재단비목"]
    ),
    "task": find_col(
        edf.columns,
        ["과제번호", "과제 번호"]
    ),
    "fund": find_col(
        edf.columns,
        ["자금", "자금코드"]
    ),
    "account_code": find_col(
        edf.columns,
        ["계정과목", "계정코드", "G/L 계정"]
    ),
    "date": find_col(
        edf.columns,
        ["전기일", "증빙일", "입력일"]
    ),
    "amount": find_col(
        edf.columns,
        ["금액(현지 통화)", "금액(현지통화)", "금액"]
    ),
    "drcr": find_col(
        edf.columns,
        ["차변/대변 지시자", "차변대변지시자"]
    )
}

missing_e = [
    k for k in [
        "dept", "nrf", "task", "fund",
        "account_code", "date", "amount"
    ]
    if not ec[k]
]

if missing_e:
    st.error(
        f"집행 시트에서 필요한 열을 찾지 못했습니다: {missing_e}. "
        "전표 상세 시트를 선택해 주세요."
    )
    st.write(list(edf.columns))
    st.stop()


# ─────────────────────────────────────────────
# 예산 데이터 정리
# ─────────────────────────────────────────────

budget = pd.DataFrame()

for k, col in bc.items():
    if col is not None:
        budget[k] = bdf[col]

for c in ["dept", "dept_code", "task", "nrf", "account", "fund"]:
    budget[c] = budget[c].fillna("").astype(str).str.strip()

budget["budget_amt"] = pd.to_numeric(
    budget["budget"],
    errors="coerce"
)

budget["task_key"] = budget["task"].map(clean_task)

budget["account_key"] = budget["account_code"].map(
    lambda x: clean_task(x) if pd.notna(x) else ""
)

budget["fund_key"] = budget["fund"].map(clean_task)
budget["dept_key"] = budget["dept"].map(norm)
budget["nrf_key"] = budget["nrf"].map(norm)

budget = budget[
    budget["budget_amt"].notna()
    & (budget["dept"] != "")
    & (budget["task_key"] != "")
    & (budget["fund_key"] != "")
].copy()

budget["영역구분"] = budget.apply(
    lambda r: area_class(r["task"], r.get("area", "")),
    axis=1
)

budget["부서명(부서코드)"] = (
    budget["dept"]
    + " ("
    + budget["dept_code"].astype(str).str.replace(
        r"\.0$", "", regex=True
    )
    + ")"
)


# ─────────────────────────────────────────────
# 집행 데이터 정리
# ─────────────────────────────────────────────

exp = pd.DataFrame()

for k, col in ec.items():
    if col is not None:
        exp[k] = edf[col]

for c in ["dept", "nrf", "task", "fund", "account_code"]:
    exp[c] = exp[c].fillna("").astype(str).str.strip()

exp["date"] = pd.to_datetime(
    exp["date"],
    errors="coerce"
)

exp["amount"] = pd.to_numeric(
    exp["amount"],
    errors="coerce"
)

exp = exp[
    exp["date"].notna()
    & exp["amount"].notna()
].copy()

exp["task_key"] = exp["task"].map(clean_task)
exp["account_key"] = exp["account_code"].map(clean_task)
exp["fund_key"] = exp["fund"].map(clean_task)
exp["dept_key"] = exp["dept"].map(norm)
exp["nrf_key"] = exp["nrf"].map(norm)

exp["영역구분"] = exp["task"].map(area_class)
exp["월"] = exp["date"].dt.to_period("M").astype(str)

# 원본 금액의 부호 유지
exp["집행액"] = exp["amount"]


# ─────────────────────────────────────────────
# 예산과 집행 연결
# ─────────────────────────────────────────────

key = [
    "dept_key",
    "task_key",
    "fund_key",
    "nrf_key",
    "account_key"
]

budget_group = budget.groupby(
    key,
    dropna=False,
    as_index=False
).agg(
    예산액=("budget_amt", "sum"),
    부서명=("dept", "first"),
    부서코드=("dept_code", "first"),
    과제번호=("task", "first"),
    자금코드=("fund", "first"),
    한국연구재단비목=("nrf", "first"),
    교내계정과목=("account", "first"),
    영역구분=("영역구분", "first")
)

exp_group = exp.groupby(
    key,
    dropna=False,
    as_index=False
).agg(
    누적집행액=("집행액", "sum")
)

joined = budget_group.merge(
    exp_group,
    on=key,
    how="left"
)

joined["누적집행액"] = joined["누적집행액"].fillna(0)
joined["잔액"] = joined["예산액"] - joined["누적집행액"]

joined["집행률(%)"] = joined.apply(
    lambda r: (
        r["누적집행액"] / r["예산액"] * 100
        if r["예산액"] else 0
    ),
    axis=1
)


# ─────────────────────────────────────────────
# 조회 조건
# ─────────────────────────────────────────────

st.sidebar.header("조회 조건")

areas = ["전체"] + sorted(
    set(budget_group["영역구분"].dropna().tolist())
    | set(exp["영역구분"].dropna().tolist())
)

sel_area = st.sidebar.selectbox("사업 영역", areas)

sel_dept = st.sidebar.multiselect(
    "부서",
    sorted(set(budget["dept"].dropna()) | set(exp["dept"].dropna()))
)

sel_task = st.sidebar.multiselect(
    "과제번호",
    sorted(set(budget["task_key"].dropna()) | set(exp["task_key"].dropna()))
)

sel_fund = st.sidebar.multiselect(
    "자금코드",
    sorted(set(budget["fund_key"].dropna()) | set(exp["fund_key"].dropna()))
)

sel_nrf = st.sidebar.multiselect(
    "한국연구재단비목",
    sorted(set(budget["nrf"].dropna()) | set(exp["nrf"].dropna()))
)


def filter_common(frame, task_col="task_key"):
    out = frame.copy()

    if sel_area != "전체":
        out = out[out["영역구분"] == sel_area]

    if sel_dept:
        if "dept" in out.columns:
            out = out[out["dept"].isin(sel_dept)]
        else:
            out = out[out["부서명"].isin(sel_dept)]

    if sel_task:
        out = out[out[task_col].isin(sel_task)]

    if sel_fund:
        out = out[out["fund_key"].isin(sel_fund)]

    if sel_nrf:
        if "nrf" in out.columns:
            out = out[out["nrf"].isin(sel_nrf)]
        else:
            out = out[out["한국연구재단비목"].isin(sel_nrf)]

    return out


bf = filter_common(budget)
ef = filter_common(exp)


# ─────────────────────────────────────────────
# 필터 적용 후 예산·집행 재집계
# ─────────────────────────────────────────────

bagg = bf.groupby(
    key,
    dropna=False,
    as_index=False
).agg(
    예산액=("budget_amt", "sum"),
    부서명=("dept", "first"),
    부서코드=("dept_code", "first"),
    과제번호=("task", "first"),
    자금코드=("fund", "first"),
    한국연구재단비목=("nrf", "first"),
    교내계정과목=("account", "first"),
    영역구분=("영역구분", "first")
)

eagg = ef.groupby(
    key,
    dropna=False,
    as_index=False
).agg(
    누적집행액=("집행액", "sum")
)

summary = bagg.merge(
    eagg,
    on=key,
    how="left"
)

summary["누적집행액"] = summary["누적집행액"].fillna(0)
summary["잔액"] = summary["예산액"] - summary["누적집행액"]

summary["집행률(%)"] = summary.apply(
    lambda r: (
        r["누적집행액"] / r["예산액"] * 100
        if r["예산액"] else 0
    ),
    axis=1
)


# ─────────────────────────────────────────────
# KPI 카드
# ─────────────────────────────────────────────

total_b = bf["budget_amt"].sum()
total_e = ef["집행액"].sum()

k1, k2, k3, k4 = st.columns(4)

k1.metric("선택 예산액", money(total_b))
k2.metric("누적 순집행액", money(total_e))
k3.metric("예산 잔액(단순 계산)", money(total_b - total_e))

k4.metric(
    "집행률",
    f"{total_e / total_b * 100:.1f}%"
    if total_b else "—"
)

st.caption(
    f"예산 시트: {bsheet} · 집행 시트: {exsheet} · "
    "집행일자: 전기일 · 금액은 원본 부호 유지(환입/취소는 음수)"
)


# ─────────────────────────────────────────────
# 1. 월별 집행액 그래프
# ─────────────────────────────────────────────

st.subheader("월별 집행액")

monthly = (
    ef.groupby("월", as_index=False)["집행액"]
    .sum()
    .sort_values("월")
)

if not monthly.empty:
    fig = px.bar(
        monthly,
        x="월",
        y="집행액",
        text="집행액",
        title="월별 순집행액",
        color_discrete_sequence=[COLOR_EXECUTION]
    )

    fig.update_traces(
        texttemplate="%{text:,.0f}원",
        textposition="outside",
        cliponaxis=False,
        hovertemplate=(
            "%{x}<br>"
            "집행액: %{y:,.0f}원"
            "<extra></extra>"
        )
    )

    fig.update_yaxes(tickformat=",.0f")

    fig.update_layout(
        xaxis_title="월",
        yaxis_title="집행액(원)",
        showlegend=False,
        title_font_color=COLOR_TITLE,
        plot_bgcolor="white",
        paper_bgcolor="white"
    )

    st.plotly_chart(fig, use_container_width=True)

    monthly["누적집행액"] = monthly["집행액"].cumsum()

    monthly_display = monthly.rename(
        columns={"월": "집행월"}
    ).copy()

    # 표에만 합계 행 추가: 그래프에 사용하는 monthly 원본은 변경하지 않음
    monthly_total = pd.DataFrame([{
        "집행월": "합계",
        "집행액": monthly_display["집행액"].sum(),
        "누적집행액": monthly_display["누적집행액"].iloc[-1]
    }])
    monthly_display = pd.concat(
        [monthly_display, monthly_total],
        ignore_index=True
    )

    monthly_display = format_money_columns(
        monthly_display,
        ["집행액", "누적집행액"]
    )

    st.dataframe(
        right_align_columns(
            monthly_display,
            ["집행액", "누적집행액"]
        ),
        use_container_width=True,
        hide_index=True
    )

else:
    st.warning(
        "선택 조건에 해당하는 유효한 전기일/금액이 없습니다."
    )


# ─────────────────────────────────────────────
# 2. 비목별 예산 구성 원그래프
# ─────────────────────────────────────────────

st.subheader("비목별 예산 구성")

nrf_source = bf.copy()

nrf_source["비목별표시"] = (
    nrf_source["nrf"]
    .replace("", "비목 미분류")
)

nrf_budget_summary = (
    nrf_source.groupby(
        "비목별표시",
        as_index=False
    )["budget_amt"]
    .sum()
    .rename(columns={"budget_amt": "예산액"})
)

nrf_execution_source = ef.copy()
nrf_execution_source["비목별표시"] = (
    nrf_execution_source["nrf"].replace("", "비목 미분류")
)
nrf_execution_summary = (
    nrf_execution_source.groupby(
        "비목별표시", as_index=False
    )["집행액"]
    .sum()
    .rename(columns={"집행액": "누적집행액"})
)

nrf_summary = nrf_budget_summary.merge(
    nrf_execution_summary,
    on="비목별표시",
    how="left"
)
nrf_summary["누적집행액"] = nrf_summary["누적집행액"].fillna(0)
nrf_summary["잔액"] = nrf_summary["예산액"] - nrf_summary["누적집행액"]
nrf_summary["집행률(%)"] = nrf_summary.apply(
    lambda r: r["누적집행액"] / r["예산액"] * 100 if r["예산액"] else 0,
    axis=1
)
nrf_summary = nrf_summary.sort_values("예산액", ascending=False)

if not nrf_summary.empty:
    # 원그래프와 표를 좌우로 배치. 합계 행은 표에만 추가하고 그래프 데이터는 유지.
    pie_col, table_col = st.columns([1, 1.25], gap="medium")

    with pie_col:
        fig = px.pie(
            nrf_summary,
            names="비목별표시",
            values="예산액",
            title="전체 예산 중 비목별 구성 비율",
            color_discrete_sequence=BLUE_PALETTE
        )

        fig.update_traces(
            textposition="inside",
            texttemplate="%{label}<br>%{percent}",
            hovertemplate=(
                "%{label}<br>"
                "예산액: %{value:,.0f}원<br>"
                "구성 비율: %{percent}"
                "<extra></extra>"
            )
        )

        fig.update_layout(
            legend_title_text="한국연구재단비목",
            title_font_color=COLOR_TITLE,
            title_font_size=16,
            paper_bgcolor="white",
            margin=dict(l=10, r=10, t=60, b=10)
        )

        st.plotly_chart(fig, use_container_width=True)

    with table_col:
        st.markdown(
            "<div style='font-size:16px; font-weight:600; color:#111D37; "
            "line-height:1.2; margin:0 0 0.8rem 0;'>비목별 예산 및 집행 현황</div>",
            unsafe_allow_html=True
        )
        nrf_display = nrf_summary[[
            "비목별표시", "예산액", "누적집행액", "잔액", "집행률(%)"
        ]].copy()
        nrf_display = nrf_display.rename(columns={"비목별표시": "한국연구재단비목"})

        nrf_total_budget = nrf_summary["예산액"].sum()
        nrf_total_execution = nrf_summary["누적집행액"].sum()
        nrf_total = pd.DataFrame([{
            "한국연구재단비목": "합계",
            "예산액": nrf_total_budget,
            "누적집행액": nrf_total_execution,
            "잔액": nrf_total_budget - nrf_total_execution,
            "집행률(%)": nrf_total_execution / nrf_total_budget * 100 if nrf_total_budget else 0
        }])
        nrf_display = pd.concat([nrf_display, nrf_total], ignore_index=True)
        nrf_display = format_money_columns(
            nrf_display,
            ["예산액", "누적집행액", "잔액"]
        )
        nrf_display["집행률(%)"] = nrf_display["집행률(%)"].map(lambda x: f"{x:.1f}%")

        st.dataframe(
            right_align_columns(
                nrf_display,
                ["예산액", "누적집행액", "잔액", "집행률(%)"]
            ),
            use_container_width=True,
            hide_index=True,
            height=min(max(180, 35 * (len(nrf_display) + 1) + 3), 20000)
        )

else:
    st.info("표시할 비목별 예산이 없습니다.")


# ─────────────────────────────────────────────
# 3. 부서별 예산 대비 집행 그래프
# ─────────────────────────────────────────────

st.subheader("부서별 예산 대비 집행")

dept_summary = (
    summary.groupby(
        "부서명",
        as_index=False
    )[["예산액", "누적집행액", "잔액"]]
    .sum()
    .sort_values("예산액", ascending=False)
)

dept_summary["집행률(%)"] = dept_summary.apply(
    lambda r: r["누적집행액"] / r["예산액"] * 100 if r["예산액"] else 0,
    axis=1
)

if not dept_summary.empty:
    # 그래프는 예산액 상위 15개 부서만 표시
    dept_chart = dept_summary.head(15).copy()

    dept_long = dept_chart.melt(
        id_vars="부서명",
        value_vars=["예산액", "누적집행액"],
        var_name="구분",
        value_name="금액"
    )

    fig = px.bar(
        dept_long,
        x="부서명",
        y="금액",
        color="구분",
        barmode="group",
        title="부서별 예산·누적집행 비교 (예산액 상위 15개)",
        color_discrete_map={
            "예산액": COLOR_BUDGET,
            "누적집행액": COLOR_EXECUTION
        }
    )

    fig.update_traces(
        hovertemplate=(
            "%{x}<br>"
            "%{fullData.name}: %{y:,.0f}원"
            "<extra></extra>"
        )
    )

    fig.update_yaxes(tickformat=",.0f")

    fig.update_layout(
        xaxis_title="부서",
        yaxis_title="금액(원)",
        xaxis_tickangle=-35,
        legend_title_text="구분",
        title_font_color=COLOR_TITLE,
        plot_bgcolor="white",
        paper_bgcolor="white"
    )

    st.plotly_chart(fig, use_container_width=True)

    dept_display = dept_summary.copy()
    dept_total_budget = dept_summary["예산액"].sum()
    dept_total_execution = dept_summary["누적집행액"].sum()
    dept_total = pd.DataFrame([{
        "부서명": "합계",
        "예산액": dept_total_budget,
        "누적집행액": dept_total_execution,
        "잔액": dept_summary["잔액"].sum(),
        "집행률(%)": dept_total_execution / dept_total_budget * 100 if dept_total_budget else 0
    }])
    dept_display = pd.concat([dept_display, dept_total], ignore_index=True)
    dept_display = format_money_columns(
        dept_display,
        ["예산액", "누적집행액", "잔액"]
    )
    dept_display["집행률(%)"] = dept_display["집행률(%)"].map(lambda x: f"{x:.1f}%")

    st.dataframe(
        right_align_columns(
            dept_display,
            ["예산액", "누적집행액", "잔액", "집행률(%)"]
        ),
        use_container_width=True,
        hide_index=True,
        height=min(max(120, 35 * (len(dept_display) + 1) + 3), 20000)
    )

else:
    st.info("표시할 부서별 예산 자료가 없습니다.")


# ─────────────────────────────────────────────
# 4. 영역별 예산 대비 집행 그래프
# ─────────────────────────────────────────────

st.subheader("영역별 예산 대비 집행")

area_summary = (
    summary.groupby(
        "영역구분",
        as_index=False
    )[["예산액", "누적집행액", "잔액"]]
    .sum()
)

area_summary["집행률(%)"] = area_summary.apply(
    lambda r: (
        r["누적집행액"] / r["예산액"] * 100
        if r["예산액"] else 0
    ),
    axis=1
)

area_display = area_summary.copy()
area_total_budget = area_summary["예산액"].sum()
area_total_execution = area_summary["누적집행액"].sum()
area_total = pd.DataFrame([{
    "영역구분": "합계",
    "예산액": area_total_budget,
    "누적집행액": area_total_execution,
    "잔액": area_summary["잔액"].sum(),
    "집행률(%)": area_total_execution / area_total_budget * 100 if area_total_budget else 0
}])
area_display = pd.concat([area_display, area_total], ignore_index=True)
area_display = format_money_columns(
    area_display,
    ["예산액", "누적집행액", "잔액"]
)
area_display["집행률(%)"] = area_display["집행률(%)"].map(lambda x: f"{x:.1f}%")

st.dataframe(
    right_align_columns(
        area_display,
        ["예산액", "누적집행액", "잔액", "집행률(%)"]
    ),
    use_container_width=True,
    hide_index=True
)

area_long = area_summary.melt(
    id_vars="영역구분",
    value_vars=["예산액", "누적집행액"],
    var_name="구분",
    value_name="금액"
)

fig = px.bar(
    area_long,
    x="영역구분",
    y="금액",
    color="구분",
    barmode="group",
    title="영역별 예산·누적집행 비교",
    color_discrete_map={
        "예산액": COLOR_BUDGET_SUB,
        "누적집행액": COLOR_EXECUTION_SUB
    }
)

fig.update_traces(
    hovertemplate=(
        "%{x}<br>"
        "%{fullData.name}: %{y:,.0f}원"
        "<extra></extra>"
    )
)

fig.update_yaxes(tickformat=",.0f")

fig.update_layout(
    xaxis_title="사업 영역",
    yaxis_title="금액(원)",
    legend_title_text="구분",
    title_font_color=COLOR_TITLE,
    plot_bgcolor="white",
    paper_bgcolor="white"
)

st.plotly_chart(fig, use_container_width=True)


# ─────────────────────────────────────────────
# 5. 예산 세부 집계
# ─────────────────────────────────────────────

st.subheader("예산 세부 집계")

show_cols = [
    "영역구분",
    "부서명",
    "부서코드",
    "과제번호",
    "자금코드",
    "한국연구재단비목",
    "교내계정과목",
    "예산액",
    "누적집행액",
    "잔액",
    "집행률(%)"
]

summary_display = (
    summary[show_cols]
    .sort_values("예산액", ascending=False)
    .copy()
)

summary_total_budget = summary["예산액"].sum()
summary_total_execution = summary["누적집행액"].sum()
summary_total_row = {col: "" for col in show_cols}
summary_total_row.update({
    "영역구분": "합계",
    "예산액": summary_total_budget,
    "누적집행액": summary_total_execution,
    "잔액": summary["잔액"].sum(),
    "집행률(%)": summary_total_execution / summary_total_budget * 100 if summary_total_budget else 0
})
summary_display = pd.concat(
    [summary_display, pd.DataFrame([summary_total_row])],
    ignore_index=True
)
summary_display = format_money_columns(
    summary_display,
    ["예산액", "누적집행액", "잔액"]
)
summary_display["집행률(%)"] = summary_display["집행률(%)"].map(
    lambda x: f"{float(x):.1f}%" if pd.notna(x) and str(x).strip() != "" else ""
)

st.dataframe(
    right_align_columns(
        summary_display,
        ["예산액", "누적집행액", "잔액"]
    ),
    use_container_width=True,
    hide_index=True
)


# ─────────────────────────────────────────────
# 6. 집행내역 상세
# ─────────────────────────────────────────────

st.subheader("집행내역 상세")

detail_cols = [
    "date",
    "dept",
    "task",
    "fund",
    "nrf",
    "account_code",
    "amount",
    "월"
]

detail = (
    ef[detail_cols]
    .rename(
        columns={
            "date": "전기일",
            "dept": "지출 부서",
            "task": "과제번호",
            "fund": "자금",
            "nrf": "연구재단 비목",
            "account_code": "계정과목",
            "amount": "원본 금액",
            "월": "집행월"
        }
    )
    .sort_values("전기일", ascending=False)
    .copy()
)

detail_total_row = {col: "" for col in detail.columns}
detail_total_row["지출 부서"] = "합계"
detail_total_row["원본 금액"] = pd.to_numeric(detail["원본 금액"], errors="coerce").sum()
detail = pd.concat([detail, pd.DataFrame([detail_total_row])], ignore_index=True)
detail = format_money_columns(
    detail,
    ["원본 금액"]
)
detail["전기일"] = pd.to_datetime(detail["전기일"], errors="coerce").dt.strftime("%Y-%m-%d").fillna("")

st.dataframe(
    right_align_columns(
        detail,
        ["원본 금액"]
    ),
    use_container_width=True,
    hide_index=True
)


# ─────────────────────────────────────────────
# 7. 집행내역 검증
# ─────────────────────────────────────────────

matched_keys = exp_group.merge(
    budget_group[key],
    on=key,
    how="inner"
)

st.subheader("집행내역 검증")

c1, c2, c3 = st.columns(3)

c1.metric(
    "집행 원본 행 수",
    f"{len(exp):,}"
)

c2.metric(
    "예산과 연결된 집행 조합",
    f"{len(matched_keys):,}"
)

c3.metric(
    "예산과 미연결 집행 조합",
    f"{len(exp_group) - len(matched_keys):,}"
)

with st.expander("미연결 집행 내역 확인"):
    unmatched = exp_group.merge(
        budget_group[key],
        on=key,
        how="left",
        indicator=True
    )

    unmatched = unmatched[
        unmatched["_merge"] == "left_only"
    ].drop(columns="_merge")

    unmatched_display = unmatched.copy()
    unmatched_total_row = {col: "" for col in unmatched_display.columns}
    if "dept_key" in unmatched_display.columns:
        unmatched_total_row["dept_key"] = "합계"
    unmatched_total_row["누적집행액"] = unmatched_display["누적집행액"].sum() if "누적집행액" in unmatched_display.columns else 0
    unmatched_display = pd.concat(
        [unmatched_display, pd.DataFrame([unmatched_total_row])],
        ignore_index=True
    )
    unmatched_display = format_money_columns(
        unmatched_display,
        ["누적집행액"]
    )

    st.dataframe(
        right_align_columns(
            unmatched_display,
            ["누적집행액"]
        ),
        use_container_width=True,
        hide_index=True
    )

    st.caption(
        "미연결은 부서명·과제번호·자금코드·연구재단비목·"
        "계정코드가 예산 파일과 정확히 일치하지 않는 경우입니다. "
        "예산과 연결되지 않은 집행액도 월별 집행 그래프에는 포함됩니다."
    )


# ─────────────────────────────────────────────
# 8. 예산 대비 집행 집계 CSV 다운로드
# ─────────────────────────────────────────────

csv = summary[show_cols].to_csv(
    index=False
).encode("utf-8-sig")

st.download_button(
    "예산 대비 집행 집계 CSV 다운로드",
    data=csv,
    file_name="예산대비집행_집계.csv",
    mime="text/csv"
)


# ─────────────────────────────────────────────
# 한국연구재단 월별제출자료
# 병합 셀 유지 + 숫자 오른쪽 정렬 + 백만원 단위 표시
# ─────────────────────────────────────────────

import html
import openpyxl
from numbers import Number

st.divider()
st.header("한국연구재단 월별제출자료")

if nrf_monthly_file is None:
    st.info(
        "왼쪽 파일 업로드 메뉴에서 "
        "'③ 한국연구재단 월별제출자료 엑셀 파일'을 업로드해 주세요."
    )

else:
    try:
        nrf_monthly_file.seek(0)

        wb = openpyxl.load_workbook(
            nrf_monthly_file,
            data_only=True
        )

        ws = wb[wb.sheetnames[0]]

        # ─────────────────────────────────────
        # 병합 셀의 시작 위치와 범위 기록
        # ─────────────────────────────────────
        merged_map = {}
        merged_anchors = {}

        for merged_range in ws.merged_cells.ranges:
            r1 = merged_range.min_row
            r2 = merged_range.max_row
            c1 = merged_range.min_col
            c2 = merged_range.max_col

            merged_anchors[(r1, c1)] = {
                "rowspan": r2 - r1 + 1,
                "colspan": c2 - c1 + 1,
            }

            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    merged_map[(r, c)] = (r1, c1)

        # ─────────────────────────────────────
        # 표 스타일
        # ─────────────────────────────────────
        table_style = (
            "border-collapse: collapse;"
            "width: 100%;"
            "table-layout: auto;"
            "font-family: inherit;"
            "font-size: 14px;"
            "color: inherit;"
        )

        # 일반 텍스트 셀: 가운데 정렬
        cell_style = (
            "border: 1px solid #DDDDDD;"
            "padding: 8px 10px;"
            "text-align: center;"
            "vertical-align: middle;"
            "white-space: nowrap;"
        )
        # 숫자 셀: 오른쪽 정렬
        number_style = (
            "border: 1px solid #DDDDDD;"
            "padding: 8px 10px;"
            "text-align: right;"
            "vertical-align: middle;"
            "white-space: nowrap;"
        )

        # 상단 헤더 셀
        header_style = (
            "border: 1px solid #DDDDDD;"
            "padding: 8px 10px;"
            "text-align: center;"
            "vertical-align: middle;"
            "font-weight: 600;"
            "background-color: #F0F2F6;"
            "color: #262730;"
            "white-space: nowrap;"
        )

        # 표 위 단위 표시
        st.markdown(
            "<div style='text-align: right; "
            "font-size: 13px; margin-bottom: 6px;'>"
            "(단위: 백만원)"
            "</div>",
            unsafe_allow_html=True
        )

        
        # ─────────────────────────────────────
        # HTML 표 시작
        # ─────────────────────────────────────
        html_parts = [
            '<div style="width:100%; overflow-x:auto;">',
            f'<table style="{table_style}">',
            "<tbody>"
        ]

        # ─────────────────────────────────────
        # 엑셀 행과 열 출력
        # ─────────────────────────────────────
        for r in range(1, ws.max_row + 1):
            html_parts.append("<tr>")

            for c in range(1, ws.max_column + 1):

                # 병합 셀 내부는 출력하지 않음
                anchor = merged_map.get((r, c))

                if anchor is not None and anchor != (r, c):
                    continue

                value = ws.cell(r, c).value

                # ─────────────────────────────
                # 셀 값 및 숫자 변환
                # ─────────────────────────────
                is_number = (
                    isinstance(value, Number)
                    and not isinstance(value, bool)
                )

                if value is None:
                    text = ""

                elif is_number:
                    # 원 단위 → 백만원 단위
                    converted_value = float(value) / 1_000_000

                    # 소수점 첫째 자리까지 표시
                    text = f"{converted_value:,.1f}"

                else:
                    text = html.escape(
                        str(value)
                    ).replace("\n", "<br>")

                # ─────────────────────────────
                # 병합 범위 적용
                # ─────────────────────────────
                merge_attrs = ""

                if (r, c) in merged_anchors:
                    info = merged_anchors[(r, c)]

                    if info["rowspan"] > 1:
                        merge_attrs += (
                            f' rowspan="{info["rowspan"]}"'
                        )

                    if info["colspan"] > 1:
                        merge_attrs += (
                            f' colspan="{info["colspan"]}"'
                        )

                # ─────────────────────────────
                # 셀 스타일 결정
                # ─────────────────────────────
                if r <= 5:
                    # 상단 5행은 헤더
                    style = header_style

                elif is_number:
                    # 숫자 셀은 오른쪽 정렬
                    style = number_style

                else:
                    # 일반 텍스트 셀은 가운데 정렬
                    style = cell_style

                # 셀 추가
                html_parts.append(
                    f'<td{merge_attrs} style="{style}">{text}</td>'
                )

            html_parts.append("</tr>")

        # ─────────────────────────────────────
        # HTML 표 종료 및 화면 표시
        # ─────────────────────────────────────
        html_parts.extend([
            "</tbody>",
            "</table>",
            "</div>"
        ])

        st.markdown(
            "".join(html_parts),
            unsafe_allow_html=True
        )

        # 원본 파일명 표시
        st.caption(
            f"원본 파일: {nrf_monthly_file.name}"
            " | 병합 셀 유지 · 숫자 오른쪽 정렬"
            " · 백만원 단위 표시"
        )

        # ─────────────────────────────────────
        # 원본 엑셀 다운로드
        # ─────────────────────────────────────
        nrf_monthly_file.seek(0)

        st.download_button(
            label="한국연구재단 월별제출자료 원본 다운로드",
            data=nrf_monthly_file.getvalue(),
            file_name=nrf_monthly_file.name,
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            key="download_nrf_monthly"
        )

    except Exception as e:
        st.error(
            "월별제출자료를 표시하는 중 오류가 발생했습니다."
        )
        st.exception(e)
